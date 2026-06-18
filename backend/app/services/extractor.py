"""
Text Extractor — Mengekstraksi teks dari berbagai format dokumen.

Format yang didukung:
  .txt, .pdf, .docx, .xlsx, .csv, .ppt, .pptx
"""

import logging
import os
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


class TextExtractor:
    """Ekstraksi teks dari file berdasarkan tipe/ekstensi."""

    _BULLET_PREFIX_RE = re.compile(r"^\s*[\u2022\u2023\u25E6\u25AA\u25AB\u25CF\u2219\-–—\*•·]+\s*")
    _LEADING_SYMBOL_RE = re.compile(r"^\s*[^\w]+")
    _FOOTER_PAGE_PREFIX_RE = re.compile(r"^\s*\d+\s+")
    _FOOTER_PAGE_SUFFIX_RE = re.compile(r"\s+\d+(\s*/\s*\d+)?\s*$")
    _FOOTER_REGEXES = [
        re.compile(r"^\s*\d+\s*$", re.IGNORECASE),
        re.compile(r"^\s*\d+\s*/\s*\d+\s*$", re.IGNORECASE),
        re.compile(r"^\s*(page|halaman)\s+\d+(\s*/\s*\d+)?\s*$", re.IGNORECASE),
        re.compile(r"^\s*©\s*\d{4}.*$", re.IGNORECASE),
        re.compile(r"^\s*copyright\b.*$", re.IGNORECASE),
        re.compile(r"^\s*all rights reserved\b.*$", re.IGNORECASE),
    ]
    _FOOTER_MAX_LEN = 80
    _FOOTER_FREQ_MAX_LEN = 300

    @staticmethod
    def _normalize_line(text: str) -> str:
        return re.sub(r"\s+", " ", (text or "").strip().lower())

    @classmethod
    def _normalize_footer_line(cls, text: str) -> str:
        normalized = cls._normalize_line(text)
        normalized = cls._FOOTER_PAGE_PREFIX_RE.sub("", normalized)
        normalized = cls._FOOTER_PAGE_SUFFIX_RE.sub("", normalized)
        return re.sub(r"\s+", " ", normalized).strip()

    @classmethod
    def _strip_bullet_prefix(cls, line: str) -> str:
        stripped = cls._BULLET_PREFIX_RE.sub("", line or "")
        return cls._LEADING_SYMBOL_RE.sub("", stripped)

    @classmethod
    def _is_footer_line(cls, line: str) -> bool:
        if not line:
            return False
        if len(line) > cls._FOOTER_MAX_LEN:
            return False
        return any(rgx.match(line) for rgx in cls._FOOTER_REGEXES)

    @classmethod
    def _build_footer_set(cls, pages_lines: list[list[str]], min_ratio: float = 0.6) -> set[str]:
        counts: Counter[str] = Counter()
        total_pages = max(len(pages_lines), 1)
        for lines in pages_lines:
            seen = set()
            for line in lines:
                normalized = cls._normalize_footer_line(line)
                if not normalized or len(normalized) > cls._FOOTER_FREQ_MAX_LEN:
                    continue
                if normalized in seen:
                    continue
                seen.add(normalized)
                counts[normalized] += 1

        threshold = max(2, int(total_pages * min_ratio))
        return {line for line, c in counts.items() if c >= threshold}

    @classmethod
    def _clean_lines(cls, lines: list[str], footer_set: set[str]) -> list[str]:
        cleaned: list[str] = []
        for line in lines:
            if not line:
                continue
            normalized = cls._normalize_footer_line(line)
            if normalized in footer_set:
                continue
            if cls._is_footer_line(line):
                continue
            stripped = cls._strip_bullet_prefix(line)
            stripped = stripped.strip()
            if stripped:
                cleaned.append(stripped)
        return cleaned

    @staticmethod
    def _sanitize_text(text: Optional[str]) -> Optional[str]:
        """Bersihkan karakter NUL yang tidak diterima PostgreSQL."""
        if text is None:
            return None
        return text.replace("\x00", "")

    def extract(self, file_path: str, file_type: str) -> Optional[str]:
        """Entry point: pilih extractor sesuai tipe file."""
        extractors = {
            ".txt":  self._extract_txt,
            ".pdf":  self._extract_pdf,
            ".docx": self._extract_docx,
            ".xlsx": self._extract_xlsx,
            ".csv":  self._extract_csv,
            ".ppt":  self._extract_ppt,
            ".pptx": self._extract_pptx,
        }

        extractor = extractors.get(file_type)
        if extractor is None:
            logger.warning(f"Tidak ada extractor untuk tipe: {file_type}")
            return None

        try:
            text = extractor(file_path)
            if text and text.strip():
                return self._sanitize_text(text).strip()
            return None
        except Exception as e:
            logger.error(f"Gagal ekstraksi {file_path}: {e}")
            return None

    def extract_from_bytes(self, content_bytes: bytes, file_type: str) -> Optional[str]:
        """Ekstraksi dari bytes (mis. file cloud) menggunakan temp file."""
        if not content_bytes:
            return None

        with tempfile.NamedTemporaryFile(delete=False, suffix=file_type) as tmp:
            tmp.write(content_bytes)
            tmp_path = tmp.name

        try:
            text = self.extract(tmp_path, file_type)
            return self._sanitize_text(text)
        finally: #selalu dijalankan meskipun ada error, untuk hapus temp file
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    # ------------------------------------------------------------------
    # .TXT
    # ------------------------------------------------------------------
    def _extract_txt(self, path: str) -> Optional[str]:
        for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        return None

    # ------------------------------------------------------------------
    # .PDF
    # ------------------------------------------------------------------
    def _extract_pdf(self, path: str) -> Optional[str]:
        from PyPDF2 import PdfReader

        reader = PdfReader(path)
        pages_lines: list[list[str]] = []
        raw_pages: list[tuple[int, list[str]]] = []
        for page_num, page in enumerate(reader.pages, 1):
            page_text = page.extract_text()
            if not page_text or not page_text.strip():
                continue
            lines = [ln for ln in page_text.splitlines() if ln.strip()]
            raw_pages.append((page_num, lines))
            pages_lines.append(lines)

        if not raw_pages:
            return None

        footer_set = self._build_footer_set(pages_lines)
        parts: list[str] = []
        for page_num, lines in raw_pages:
            cleaned_lines = self._clean_lines(lines, footer_set)
            if not cleaned_lines:
                continue
            parts.append(f"[Page {page_num}]")
            parts.append("\n".join(cleaned_lines))

        return "\n".join(parts) if parts else None

    # ------------------------------------------------------------------
    # .DOCX
    # ------------------------------------------------------------------
    def _extract_docx(self, path: str) -> Optional[str]:
        from docx import Document

        doc = Document(path)
        parts = []

        # Paragraf
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        # Tabel
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts) if parts else None

    # ------------------------------------------------------------------
    # .XLSX
    # ------------------------------------------------------------------
    def _extract_xlsx(self, path: str) -> Optional[str]:
        import pandas as pd

        parts = []
        xls = pd.ExcelFile(path)
        for sheet in xls.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet)
            parts.append(f"[Sheet: {sheet}]")
            parts.append(df.to_string(index=False))
        return "\n".join(parts) if parts else None

    # ------------------------------------------------------------------
    # .CSV
    # ------------------------------------------------------------------
    def _extract_csv(self, path: str) -> Optional[str]:
        import pandas as pd

        try:
            df = pd.read_csv(path)
        except Exception:
            try:
                df = pd.read_csv(path, sep=";")
            except Exception:
                return None
        return df.to_string(index=False)

    # ------------------------------------------------------------------
    # .PPTX
    # ------------------------------------------------------------------
    def _extract_pptx(self, path: str) -> Optional[str]:
        from pptx import Presentation

        prs = Presentation(path)
        slides_lines: list[list[str]] = []
        raw_slides: list[tuple[int, list[str]]] = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_lines: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    for line in str(shape.text).splitlines():
                        if line.strip():
                            slide_lines.append(line)
            if slide_lines:
                raw_slides.append((slide_num, slide_lines))
                slides_lines.append(slide_lines)

        if not raw_slides:
            return None

        footer_set = self._build_footer_set(slides_lines)
        parts: list[str] = []
        for slide_num, lines in raw_slides:
            cleaned_lines = self._clean_lines(lines, footer_set)
            if not cleaned_lines:
                continue
            parts.append(f"[Slide {slide_num}]")
            parts.extend(cleaned_lines)
        return "\n".join(parts) if parts else None

    def _convert_ppt_to_pptx(self, path: str) -> Optional[str]:
        """Konversi legacy .ppt ke .pptx via LibreOffice soffice."""
        if not path.lower().endswith(".ppt"):
            return None

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            logger.error(f"soffice/libreoffice tidak ditemukan di PATH: {path}")
            return None

        if not os.path.exists(path):
            logger.error(f"File .ppt tidak ada: {path}")
            return None

        output_dir = tempfile.mkdtemp(prefix="ppt_convert_")
        try:
            logger.info(f"Konversi .ppt → .pptx: {path}")
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pptx", "--outdir", output_dir, path],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
                timeout=60,
            )
        except subprocess.TimeoutExpired:
            logger.error(f"Timeout konversi .ppt (60s): {path}")
            return None
        except subprocess.CalledProcessError as e:
            stderr_msg = e.stderr.decode('utf-8', errors='ignore') if e.stderr else "N/A"
            logger.error(f"soffice error: {path} | {stderr_msg}")
            return None
        except Exception as e:
            logger.error(f"Exception konversi .ppt: {path} | {type(e).__name__}: {e}")
            return None

        stem = Path(path).stem
        converted = Path(output_dir) / f"{stem}.pptx"
        if not converted.exists():
            logger.error(f"Hasil konversi .pptx tidak ditemukan: {converted}")
            return None

        logger.info(f"Konversi .ppt sukses: {converted}")
        return str(converted)

    def _extract_ppt(self, path: str) -> Optional[str]:
        """Ekstrak legacy .ppt dengan konversi sementara ke .pptx."""
        converted_path = self._convert_ppt_to_pptx(path)
        if not converted_path:
            logger.warning(f"Gagal konversi .ppt: {path}")
            return None

        try:
            result = self._extract_pptx(converted_path)
            logger.info(f"Ekstrak .ppt sukses: {path}")
            return result
        except Exception as e:
            logger.error(f"Gagal ekstrak .pptx konversi: {path} | {e}")
            return None
        finally:
            try:
                os.remove(converted_path)
            except OSError:
                pass
